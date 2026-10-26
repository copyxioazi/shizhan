"""生成 4 类示例合同：PDF、DOCX、扫描 PNG、缺失附件。

生成过程必须确定性：固定内容、排版、元数据与 ZIP 时间戳。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
import zipfile
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.config import settings
from app.errors import ToolError

FIXED_ZIP_TIME = (2026, 1, 1, 0, 0, 0)
FIXED_PROPS_TIME = datetime(2026, 1, 1, 0, 0, 0)

FONT_CANDIDATES = [
    os.getenv("CN_FONT_PATH", ""),
    r"C:\Windows\Fonts\simsun.ttc",
    r"C:\Windows\Fonts\msyh.ttc",
    r"C:\Windows\Fonts\simhei.ttf",
    # Linux / Docker 常用中文字体
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
    "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
]

PURCHASE_TITLE = "设备采购合同"
PURCHASE_LINES = [
    "合同编号：CG-2026-001",
    "甲方（买方）：杭州星海科技有限公司",
    "乙方（卖方）：苏州远方设备制造有限公司",
    "合同金额：人民币 1,200,000 元（含税）",
    "生效日期：2026年3月1日",
    "到期日期：2027年2月28日",
    "第一条 采购标的",
    "甲方向乙方采购自动化生产设备一套，具体型号、数量与技术参数以附件《报价单》为准。",
    "乙方应保证所供设备为全新、未使用过的合格产品，并符合国家相关质量标准。",
    "第二条 付款条款",
    "合同签订后 5 个工作日内，甲方向乙方支付预付款 60%。",
    "剩余款项于设备验收合格后 90 天内支付。",
    "乙方应在收款前向甲方开具等额增值税专用发票。",
    "第三条 交付条款",
    "乙方应于 2026 年 5 月 31 日前将全部设备交付至甲方指定地点，运输费用由乙方承担。",
    "设备交付时，双方应共同清点数量并签署到货确认单。",
    "第四条 验收条款",
    "验收标准：设备连续稳定运行 7 日且各项技术参数达到约定指标，视为验收合格。",
    "验收方式：甲方组织验收，乙方应派员配合；验收不合格的，乙方应在 15 日内完成整改。",
    "第五条 违约责任",
    "任何一方违反本合同约定，应向守约方支付合同总金额 10% 的违约金，并赔偿由此造成的直接损失。",
    "乙方逾期交付的，每逾期一日按合同总金额的千分之一支付违约金。",
    "第六条 保密条款",
    "双方对在履行本合同过程中知悉的对方商业秘密和保密信息负有保密义务，未经对方书面同意不得向任何第三方披露。",
    "保密义务在本合同终止后继续有效三年。",
    "第七条 数据处理与个人信息保护",
    "乙方因履行本合同处理甲方数据时，应遵守数据安全法律法规，采取必要的安全保护措施，不得超范围使用数据或泄露个人信息。",
    "第八条 知识产权",
    "本合同履行过程中形成的知识产权归甲方所有；乙方保证其交付物不侵犯任何第三方的知识产权。",
    "第九条 争议解决",
    "因本合同发生争议，双方应友好协商；协商不成的，由乙方所在地人民法院管辖。",
    "第十条 自动续约与合同期限",
    "本合同期满前 30 日，双方均未提出书面异议的，本合同自动续约一年。",
    "本合同自双方盖章之日起生效，有效期至 2027 年 2 月 28 日。",
    "第十一条 其他",
    "本合同一式两份，甲乙双方各执一份，具有同等法律效力。",
]

QUOTATION_TITLE = "报价单"
QUOTATION_LINES = [
    "报价编号：BJ-2026-001",
    "供应商：苏州远方设备制造有限公司",
    "采购方：杭州星海科技有限公司",
    "报价金额：人民币 1,200,000 元（含税）",
    "报价日期：2026年2月20日",
    "设备名称：自动化生产设备",
    "数量：1 套",
    "单价：人民币 1,200,000 元",
    "交付周期：合同生效后 90 日内",
    "质保期：验收合格后 12 个月",
    "备注：本报价单仅用于采购审批，最终以正式合同为准。",
]

SERVICE_TITLE = "技术服务合同"
SERVICE_LINES = [
    "合同编号：FW-2026-002",
    "甲方：杭州星海科技有限公司",
    "乙方：南京云图信息技术有限公司",
    "合同金额：人民币 480,000 元（含税）",
    "生效日期：2026年4月1日",
    "到期日期：2027年3月31日",
    "一、服务内容",
    "乙方向甲方提供数据平台建设与运维技术服务，包括需求分析、系统设计、实施部署和技术支持。",
    "乙方应指派具备相应资质的项目人员，并保证服务团队稳定。",
    "二、付款方式",
    "合同签订后 5 个工作日内，甲方向乙方支付预付款 20%。",
    "服务成果验收合格后 75 天内，甲方支付剩余服务费用。",
    "三、服务交付",
    "乙方应按项目计划分阶段交付需求文档、设计文档、测试报告及上线系统。",
    "四、验收标准",
    "验收标准：系统功能满足需求说明书约定，连续稳定运行 14 日且无重大缺陷，视为验收合格。",
    "验收方式：甲方组织业务部门进行验收测试，乙方应提交完整的技术文档。",
    "五、违约责任",
    "任何一方违约，应向守约方支付合同总金额 10% 的违约金，并赔偿由此造成的直接损失。",
    "六、知识产权",
    "本项目形成的知识产权归甲方所有，乙方享有在项目范围内使用的权利。",
    "七、数据保护",
    "乙方处理甲方数据应遵守数据安全和个人信息保护相关法律法规，不得擅自复制、传播或用于其他目的。",
    "八、争议解决",
    "因本合同发生争议，双方应友好协商；协商不成的，由甲方所在地人民法院管辖。",
    "九、自动续约",
    "本合同期满前 30 日，双方均未提出书面异议的，本合同自动续约一年。",
    "十、合同期限",
    "本合同自双方盖章之日起生效，有效期至 2027 年 3 月 31 日。",
    "十一、其他",
    "本合同一式两份，甲乙双方各执一份。",
]

SCAN_TITLE = "保密协议"
SCAN_LINES = [
    "甲方：杭州星海科技有限公司",
    "乙方：上海青禾咨询有限公司",
    "合同金额：人民币 50,000 元",
    "生效日期：2026年5月1日",
    "到期日期：2028年4月30日",
    "一、保密义务",
    "双方应对在合作过程中知悉的保密信息承担保密义务，未经对方书面同意不得向第三方披露。",
    "二、数据安全",
    "双方应采取措施保护个人信息和数据安全，不得超范围使用。",
    "三、费用与支付",
    "本协议不涉及费用支付；如产生费用，双方另行书面约定。",
    "四、交付与验收",
    "本协议自双方书面确认视为交付完成及验收合格。",
    "五、知识产权",
    "合作前已存在的知识产权归原权利人所有，合作成果的知识产权归甲方所有。",
    "六、验收标准",
    "本协议以双方书面确认视为验收合格。",
    "七、违约责任",
    "违约方应支付违约金 50,000 元，并赔偿由此造成的直接损失。",
    "八、争议解决",
    "争议由甲方所在地人民法院管辖。",
    "九、其他",
    "本协议一式两份，双方各执一份。",
]

OPS_TITLE = "系统运维合同"
OPS_LINES = [
    "合同编号：YW-2026-004",
    "甲方：杭州星海科技有限公司",
    "乙方：杭州衡远网络科技有限公司",
    "合同金额：人民币 360,000 元（含税）",
    "生效日期：2026年6月1日",
    "到期日期：2027年5月31日",
    "一、服务内容",
    "乙方为甲方提供机房设备与业务系统的日常运维、巡检和故障处理服务。",
    "二、付款方式",
    "合同签订后，甲方向乙方支付预付款 50%。",
    "剩余款项在服务期满后 30 天内支付。",
    "三、服务交付",
    "乙方应按月提交巡检报告、故障处理记录和运维服务总结。",
    "四、验收标准",
    "运维服务按月考核，验收合格后支付当期服务费用。",
    "五、违约责任",
    "乙方未按约定提供服务的，应支付合同总金额 5% 的违约金，并赔偿由此造成的直接损失。",
    "六、保密",
    "双方应对运维过程中知悉的保密信息承担保密义务。",
    "七、数据保护",
    "乙方处理甲方数据应遵守数据安全和个人信息保护规定。",
    "八、知识产权",
    "既有知识产权归各自所有，运维成果的知识产权归甲方所有。",
    "九、争议解决",
    "协商不成的，由乙方所在地人民法院管辖。",
    "十、自动续约",
    "合同期满前 30 日，双方未提出异议的，自动续约一年。",
    "十一、合同期限",
    "本合同有效期至 2027 年 5 月 31 日。",
]


def _find_font() -> Path:
    for candidate in FONT_CANDIDATES:
        if candidate and Path(candidate).is_file():
            return Path(candidate)
    raise ToolError(
        "未找到中文字体，可通过 CN_FONT_PATH 指定 simsun.ttc/msyh.ttc/simhei.ttf",
        code="CN_FONT_MISSING",
    )


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_pdf(path: Path, title: str, lines: list[str], font: Path) -> None:
    from reportlab import rl_config

    rl_config.invariant = 1

    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    font_name = "SampleCN"
    if font.suffix.lower() == ".ttc":
        pdfmetrics.registerFont(TTFont(font_name, str(font), subfontIndex=0))
    else:
        pdfmetrics.registerFont(TTFont(font_name, str(font)))

    body = ParagraphStyle(
        "SampleBody",
        fontName=font_name,
        fontSize=11,
        leading=22,
        spaceAfter=4,
    )
    heading = ParagraphStyle(
        "SampleHeading",
        fontName=font_name,
        fontSize=16,
        leading=26,
        alignment=TA_CENTER,
        spaceAfter=18,
    )

    document = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=56,
        rightMargin=56,
        topMargin=56,
        bottomMargin=56,
        title=title,
        author="合同审批审查系统",
        creator="generate_samples.py",
    )
    story: list = [Paragraph(title, heading), Spacer(1, 8)]
    for index, line in enumerate(lines):
        story.append(Paragraph(line, body))
        # 固定分页点，避免排版引擎差异导致页数漂移。
        if index == len(lines) // 2:
            story.append(PageBreak())
    document.build(story)


def _set_east_asian_font(run, font_name: str = "宋体") -> None:
    from docx.oxml.ns import qn

    run.font.name = font_name
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), font_name)


def _write_docx(path: Path, title: str, lines: list[str]) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    document = Document()
    from docx.oxml.ns import qn

    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(11)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(16)
    _set_east_asian_font(run)

    for line in lines:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(line)
        run.font.size = Pt(11)
        _set_east_asian_font(run)

    props = document.core_properties
    props.title = title
    props.author = "合同审批审查系统"
    props.last_modified_by = "合同审批审查系统"
    props.created = FIXED_PROPS_TIME
    props.modified = FIXED_PROPS_TIME
    props.revision = 1
    document.save(str(path))
    _normalize_docx_zip(path)


def _normalize_docx_zip(path: Path) -> None:
    """排序 ZIP 条目并固定时间戳，保证 DOCX 字节级可复现。"""
    with zipfile.ZipFile(path, "r") as source:
        items = [(info.filename, source.read(info.filename)) for info in source.infolist()]
    items.sort(key=lambda item: item[0])
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, data in items:
            info = zipfile.ZipInfo(name, date_time=FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, data)


def _wrap_text(text: str, font, max_width: float) -> list[str]:
    lines: list[str] = []
    current = ""
    for char in text:
        candidate = current + char
        if font.getlength(candidate) <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = char
    if current:
        lines.append(current)
    return lines


def _write_scan_png(path: Path, title: str, lines: list[str], font_path: Path) -> None:
    from PIL import Image, ImageDraw, ImageFont

    width, height = 1654, 2339  # A4 @ 200 DPI
    margin = 110
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.truetype(str(font_path), 44)
    body_font = ImageFont.truetype(str(font_path), 34)

    cursor_y = margin
    title_width = title_font.getlength(title)
    draw.text(
        ((width - title_width) / 2, cursor_y),
        title,
        fill="black",
        font=title_font,
    )
    cursor_y += 90

    for line in lines:
        wrapped = _wrap_text(line, body_font, width - margin * 2)
        for visual_line in wrapped:
            if cursor_y + 50 > height - margin:
                raise ToolError(
                    "扫描件内容超出单页，请缩短示例文本",
                    code="SAMPLE_OVERFLOW",
                )
            draw.text((margin, cursor_y), visual_line, fill="black", font=body_font)
            cursor_y += 52
        cursor_y += 8

    image.save(path, format="PNG", optimize=True)


def _expected_manifest() -> list[dict]:
    return [
        {
            "file": "apv_2026_001_purchase.pdf",
            "approval_code": "APV-2026-001",
            "role": "主合同",
            "format": "pdf",
            "expected_parse_status": "success",
            "expected_overall_risk": "high",
            "expected_hit_rules": ["R001", "R002", "R003", "R005"],
        },
        {
            "file": "apv_2026_001_quotation.pdf",
            "approval_code": "APV-2026-001",
            "role": "次附件",
            "format": "pdf",
            "expected_parse_status": "not_parsed",
            "expected_overall_risk": None,
            "expected_hit_rules": [],
        },
        {
            "file": "apv_2026_002_service.docx",
            "approval_code": "APV-2026-002",
            "role": "主合同",
            "format": "docx",
            "expected_parse_status": "partial",
            "expected_overall_risk": "medium",
            "expected_hit_rules": ["R002", "R003", "R008"],
            "note": "保密条款缺失导致 partial，但不阻塞规则审查",
        },
        {
            "file": "apv_2026_003_scan.png",
            "approval_code": "APV-2026-003",
            "role": "扫描件",
            "format": "png",
            "expected_parse_status": "partial",
            "expected_overall_risk": "low",
            "expected_hit_rules": [],
        },
        {
            "file": "apv_2026_004_hidden.docx",
            "approval_code": "APV-2026-004",
            "role": "缺失附件",
            "format": "docx",
            "expected_parse_status": "success",
            "expected_overall_risk": "high",
            "expected_hit_rules": ["R001", "R003", "R005"],
            "note": "Mock 中 available=0，restore 后可下载",
        },
    ]


def generate_samples(output_dir: Path | None = None, verify_ocr: bool = False) -> dict:
    """确定性生成样例并更新 manifest.json。"""
    target = Path(output_dir) if output_dir is not None else settings.sample_dir
    target.mkdir(parents=True, exist_ok=True)
    font = _find_font()

    _write_pdf(target / "apv_2026_001_purchase.pdf", PURCHASE_TITLE, PURCHASE_LINES, font)
    _write_pdf(target / "apv_2026_001_quotation.pdf", QUOTATION_TITLE, QUOTATION_LINES, font)
    _write_docx(target / "apv_2026_002_service.docx", SERVICE_TITLE, SERVICE_LINES)
    _write_scan_png(target / "apv_2026_003_scan.png", SCAN_TITLE, SCAN_LINES, font)
    _write_docx(target / "apv_2026_004_hidden.docx", OPS_TITLE, OPS_LINES)

    files = _expected_manifest()
    for item in files:
        path = target / item["file"]
        if not path.is_file():
            raise ToolError(f"示例文件生成失败：{path}", code="SAMPLE_GENERATION_FAILED")
        item["sha256"] = _sha256(path)
        item["file_size"] = path.stat().st_size

    manifest = {
        "version": 2,
        "generated": True,
        "generated_by": "scripts/generate_samples.py",
        "files": files,
    }
    (target / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    if verify_ocr:
        from app.errors import NotImplementedFeature
        from app.services.parser.ocr import ocr_image

        try:
            rows = ocr_image(target / "apv_2026_003_scan.png")
        except NotImplementedFeature as exc:
            raise ToolError(
                "OCR 尚未接入，无法执行 --verify-ocr",
                code="OCR_NOT_AVAILABLE",
                reason=str(exc),
            ) from exc
        text = "".join(row["text"] for row in rows)
        for keyword in ("保密协议", "甲方", "违约责任", "验收"):
            if keyword not in text:
                raise ToolError(
                    f"OCR 未识别出关键标签：{keyword}",
                    code="OCR_VERIFY_FAILED",
                )

    return {
        "output_dir": str(target),
        "files": [
            {
                "file": item["file"],
                "sha256": item["sha256"],
                "file_size": item["file_size"],
            }
            for item in files
        ],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="生成合同审批审查示例文件")
    parser.add_argument("--output", type=Path, default=None, help="输出目录")
    parser.add_argument("--verify-ocr", action="store_true", help="生成后用 OCR 验证关键标签")
    args = parser.parse_args()
    result = generate_samples(args.output, args.verify_ocr)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
